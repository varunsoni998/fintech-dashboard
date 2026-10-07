#!/usr/bin/env python3
"""
costing_sheet.py — writes a CustomHolidays costing sheet in the team's standard layout
(the "Final Sheet" format: Status | Notes | Booked for | BID | Description | Source | Rate | ... | B.com compare).

Usage: python costing_sheet.py trip.json output.xlsx

trip.json:
{ "sheet_name": "07Oct",
  "trip": {"name": "Mr & Mrs Mittal", "rooms": "1 Room", "pax": "2 Adults", "notes": "05th Oct to 12th Oct",
           "costing_currency": "CHF", "rates": {"USD": 94.3, "EUR": 108.2, "CHF": 118.55},
           "markup": 0.15, "gst": 0.05, "tcs": 0.02},
  "sections": [ {"title": "Zermatt 05th Oct to 07th Oct 2N",
                 "rows": [ {"status": "Blocked", "notes": "", "booked_for": "", "bid": "base",
                            "description": "Schweizerhof Zermatt 4* x1 DBL room 26-28sqm",
                            "source": "Ottila", "currency": "INR", "amount": 130499,      # or "350*2"
                            "markup": 0.2, "gst_other": null, "meal": "BB", "cancel": "2026-10-01" | "NRF",
                            "due": null, "remark": "Includes spa", "pkg": ["A"],
                            "compare": {"source": "B.com", "currency": "INR", "amount": "157389+2443",
                                        "meal": "BB", "cancel": "NRF"} } ] } ],
  "land": {"title": "Land package", "source": "weshare", "markup": 0.15, "pkg": ["A"],
           "rows": [ {"date": "2026-10-05", "description": "08 days Swiss Travel Pass", "source": "",
                      "currency": "CHF", "amount": "419*2"} ] },
  "packages": [ {"key": "A", "name": "Per couple"} ] }
amount = cost of ONE room for the whole stay (as the team enters it in column G).
If output.xlsx already exists, the sheet is added as a new tab in front (older tabs are kept).
"""
import json, sys, datetime, re, os
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, Color
from openpyxl.utils import get_column_letter as CL

# ── look of the original template (Office 2007 theme colours) ───────────────────
def th(i, tint=0.0): return Color(theme=i, tint=tint)
PINK   = PatternFill("solid", fgColor=th(5, 0.7999816888943144))   # header block (accent2 light)
BLUE   = PatternFill("solid", fgColor=th(4, 0.7999816888943144))   # column headers (accent1 light)
BEIGE  = PatternFill("solid", fgColor=th(2))                       # summary block (lt2)
BLACK  = PatternFill("solid", fgColor=th(1))
YELLOW = PatternFill("solid", fgColor="FFFFFF00")
NAVYTXT = "FF002060"
def F(name="Calibri", b=False, color=None, i=False):
    return Font(name=name, size=11, bold=b, italic=i, color=color if isinstance(color, (Color, type(None))) else color)
THIN, MED = Side(style="thin"), Side(style="medium")
def border(l=None, r=None, t=None, b=None): return Border(left=l, right=r, top=t, bottom=b)
ALLTHIN, ALLMED, LR = border(THIN, THIN, THIN, THIN), border(MED, MED, MED, MED), border(THIN, THIN)
CENTER = Alignment(horizontal="center", wrap_text=True)
FMT = {
  "USD": '_-[$$-409]* #,##0_ ;_-[$$-409]* \\-#,##0\\ ;_-[$$-409]* "-"??_ ;_-@_ ',
  "CHF": '_-* #,##0\\ [$CHF-100C]_-;\\-* #,##0\\ [$CHF-100C]_-;_-* "-"??\\ [$CHF-100C]_-;_-@_-',
  "EUR": '_-[$€-2]\\ * #,##0_-;\\-[$€-2]\\ * #,##0_-;_-[$€-2]\\ * "-"??_-;_-@_-',
  "GBP": '_-[$£-809]* #,##0_-;\\-[$£-809]* #,##0_-;_-[$£-809]* "-"??_-;_-@_-',
  "INR": '_ [$₹-4009]\\ * #,##0_ ;_ [$₹-4009]\\ * \\-#,##0_ ;_ [$₹-4009]\\ * "-"??_ ;_ @_ ',
}
ACC = '_ * #,##0.00_ ;_ * \\-#,##0.00_ ;_ * "-"??_ ;_ @_ '
WIDTHS = {"A": 9, "B": 9, "C": 13.7, "D": 14.3, "E": 44.6, "F": 13, "G": 14.7, "H": 14.6, "I": 13.5, "J": 12,
          "K": 5.6, "L": 10.5, "M": 5.8, "N": 0.3, "O": 9, "P": 14.6, "Q": 13.5, "R": 6.6, "S": 11.8, "T": 34}

def ccy_fmt(c): return FMT.get(c, f'#,##0 "{c}"')
def as_date(v):
    if isinstance(v, str) and re.fullmatch(r"\d{4}-\d{2}-\d{2}", v): return datetime.date.fromisoformat(v)
    return v
def expr(a):
    """number or arithmetic text ('350*2', '157389+2443') → formula body"""
    if a is None or a == "": return None
    if isinstance(a, (int, float)): return repr(round(a, 4)).rstrip("0").rstrip(".") if isinstance(a, float) else str(a)
    s = str(a).strip().lstrip("=").replace(",", "")
    if not re.fullmatch(r"[0-9.+\-*/() %]+", s): raise ValueError(f"amount must be a number or simple sum, got {a!r}")
    return s

def build(data, out):
    T = data["trip"]; cc = T.get("costing_currency", "USD")
    rates = dict(T.get("rates") or {})
    roe_list = list(dict.fromkeys(["USD", "EUR"] + ([cc] if cc != "INR" else []) +
                    [r.get("currency") for s in data.get("sections", []) for r in s["rows"]] +
                    [ (r.get("compare") or {}).get("currency") for s in data.get("sections", []) for r in s["rows"]] +
                    [r.get("currency") for r in (data.get("land") or {}).get("rows", [])]))
    roe_list = [c for c in roe_list if c and c != "INR"]
    pkgs = data.get("packages") or [{"key": "A", "name": "Per couple"}]
    money = ccy_fmt(cc)

    name = (data.get("sheet_name") or datetime.date.today().strftime("%d%b"))[:31]
    if os.path.exists(out):                       # same trip file → add a new dated tab in front, keep old tabs
        wb = load_workbook(out)
        base, n = name, 2
        while name in wb.sheetnames: name = f"{base} ({n})"; n += 1
        ws = wb.create_sheet(name, 0); wb.active = 0
    else:
        wb = Workbook(); ws = wb.active; ws.title = name
    ws.sheet_view.zoomScale = 81
    for k, w in WIDTHS.items(): ws.column_dimensions[k].width = w

    # ── rows 2–4: trip header block ──────────────────────────────────────────
    def block(r, label, v1, v2=None):
        ws.merge_cells(f"A{r}:C{r}")
        a = ws[f"A{r}"]; a.value = label; a.font = F(b=True, color=th(1)); a.fill = PINK; a.alignment = Alignment(horizontal="center")
        if v2 is None:
            ws.merge_cells(f"D{r}:L{r}"); spans = [("D", "L", v1)]
        else:
            ws.merge_cells(f"D{r}:F{r}"); ws.merge_cells(f"G{r}:L{r}"); spans = [("D", "F", v1), ("G", "L", v2)]
        for c0, c1, v in spans:
            c = ws[f"{c0}{r}"]; c.value = v; c.font = F(b=True, color=th(1)); c.fill = PINK
            c.alignment = Alignment(horizontal="left" if c0 == "D" else "center")
        for col in range(1, 13):
            ws.cell(r, col).border = ALLTHIN
        for col in "MN": ws[f"{col}{r}"].fill = PINK; ws[f"{col}{r}"].font = F(b=True)
        for col in "PQ": ws[f"{col}{r}"].fill = PINK
    block(2, "Trip Name", T.get("name"))
    block(3, "No of Pax/ Rooms Required", T.get("rooms"), T.get("pax"))
    block(4, "Special Notes", T.get("notes"))

    # ── rows 6–7: column headers ────────────────────────────────────────────
    for rng in ("A6:A7", "C6:C7", "D6:D7", "E6:E7", "F6:L6"): ws.merge_cells(rng)
    hdr6 = {"A": "Status ", "C": "Booked for Which Date ", "D": "BID", "E": "Description ", "F": "Costing From Supplier / Online Portals "}
    hdr7 = {"B": "Notes", "F": "Source", "G": "Rate Per Room", "H": "Cost to us", "I": "GST & Other charges",
            "J": "Total Rate  (With GST)", "K": "Meal Plan", "L": "Cancel. Policy", "M": "Due date",
            "O": "Source", "P": "Cost to us", "Q": "Difference", "R": "Meal Plan", "S": "Cancel. Policy"}
    for col in "ABCDE":
        for r in (6, 7):
            c = ws[f"{col}{r}"]; c.fill = BLUE; c.font = F("Arial", True, NAVYTXT); c.alignment = CENTER; c.border = ALLMED
    for k, v in hdr6.items(): ws[f"{k}6"] = v
    ws["F6"].font = F("Arial", True); ws["F6"].fill = BLUE; ws["F6"].alignment = CENTER
    for col in "FGHIJKL": ws[f"{col}6"].border = border(t=THIN, b=THIN, l=THIN if col == "F" else None, r=THIN if col == "L" else None)
    ws["M6"].fill = BLUE; ws["M6"].border = ALLTHIN
    ws["N6"].fill = BLACK; ws["N7"].fill = BLACK
    for col in "PQ": ws[f"{col}6"].fill = YELLOW; ws[f"{col}6"].border = ALLTHIN
    for k, v in hdr7.items():
        c = ws[f"{k}7"]; c.value = v; c.alignment = CENTER; c.border = ALLMED
        c.font = F("Arial" if k != "Q" else "Calibri", True, NAVYTXT if k == "B" else None)
        c.fill = YELLOW if k in "OPQRS" else BLUE
    ws.row_dimensions[7].height = 28.8
    ws.freeze_panes = "A8"

    def frame(r, bold=False):
        for col in range(1, 20):
            c = ws.cell(r, col)
            if col == 14: continue
            c.border = LR
            c.font = F(b=bold)

    # pass 1 — work out where the ROE block lands so row formulas can point at it
    r = 9
    layout = []
    for s in data.get("sections", []):
        layout.append(("title", r, s)); r += 2
        for row in s["rows"]: layout.append(("row", r, row)); r += 1
        r += 2
    land = data.get("land")
    if land and land.get("rows"):
        layout.append(("landtitle", r, land)); r += 2
        lr0 = r
        for row in land["rows"]: layout.append(("landrow", r, row)); r += 1
        layout.append(("landtotal", r, (lr0, r - 1))); land_total_row = r; r += 3
    else:
        land_total_row = None
    roe_row = {}
    roe_start = r + 1
    for i, c in enumerate(roe_list): roe_row[c] = roe_start + i
    ROE = lambda c: f"$F${roe_row[c]}"

    def conv(amount, currency):
        e = expr(amount)
        if e is None: return None
        cur = currency or cc
        if cur == cc: return "=" + e if not re.fullmatch(r"[0-9.]+", e) else float(e)
        if cur == "INR": return f"=({e})/{ROE(cc)}" if re.search(r"[+\-*/]", e) else f"={e}/{ROE(cc)}"
        return f"=({e})*{ROE(cur)}/{ROE(cc)}"

    selected = {p["key"]: [] for p in pkgs}
    # pass 2 — write
    for kind, r, obj in layout:
        if kind == "title":
            frame(r); c = ws[f"E{r}"]; c.value = obj["title"]; c.font = F(b=True); c.fill = YELLOW; c.alignment = Alignment(horizontal="center")
            frame(r + 1)
        elif kind == "row":
            row = obj; strong = (row.get("status") or "").lower() in ("blocked", "booked")
            frame(r, strong)
            ws[f"A{r}"] = row.get("status"); ws[f"B{r}"] = row.get("notes"); ws[f"C{r}"] = row.get("booked_for")
            ws[f"D{r}"] = row.get("bid"); ws[f"E{r}"] = row.get("description"); ws[f"F{r}"] = row.get("source")
            g = conv(row.get("amount"), row.get("currency")); ws[f"G{r}"] = g
            mk = row.get("markup")
            if mk is None and row.get("pkg"): mk = T.get("markup", 0.15)
            if mk is not None and g is not None: ws[f"H{r}"] = f"=G{r}*{round(mk*100, 2):g}%"
            if row.get("gst_other") not in (None, ""): ws[f"I{r}"] = conv(row["gst_other"], row.get("gst_other_currency") or cc)
            if g is not None and (mk is not None or row.get("gst_other") not in (None, "")):
                ws[f"J{r}"] = f"=G{r}+H{r}+I{r}"
            ws[f"K{r}"] = row.get("meal"); ws[f"L{r}"] = as_date(row.get("cancel")); ws[f"M{r}"] = as_date(row.get("due"))
            cmp_ = row.get("compare") or {}
            if cmp_:
                ws[f"O{r}"] = cmp_.get("source") or "B.com"
                p = conv(cmp_.get("amount"), cmp_.get("currency"))
                ws[f"P{r}"] = p if p is not None else cmp_.get("text")
                if p is not None and g is not None: ws[f"Q{r}"] = f"=P{r}-G{r}"
                ws[f"R{r}"] = cmp_.get("meal"); ws[f"S{r}"] = as_date(cmp_.get("cancel"))
            if row.get("remark"): ws[f"T{r}"] = row["remark"]
            for col in "GHIJPQ": ws[f"{col}{r}"].number_format = money
            for col in "LMS": ws[f"{col}{r}"].number_format = "d-mmm"
            for col in "LS": ws[f"{col}{r}"].alignment = Alignment(horizontal="center")
            if strong:
                for col in "AB":
                    if ws[f"{col}{r}"].value: ws[f"{col}{r}"].fill = YELLOW
                if ws[f"J{r}"].value: ws[f"J{r}"].font = F(b=True, color="FFFF0000")
            for k in row.get("pkg") or []:
                if k in selected: selected[k].append(r)
        elif kind == "landtitle":
            frame(r); c = ws[f"E{r}"]; c.value = obj.get("title") or "Land package"; c.font = F(b=True); c.fill = YELLOW
            ws[f"F{r}"] = obj.get("source"); frame(r + 1)
        elif kind == "landrow":
            frame(r); row = obj
            ws[f"D{r}"] = as_date(row.get("date")); ws[f"D{r}"].number_format = "d-mmm" if isinstance(ws[f"D{r}"].value, datetime.date) else "General"
            ws[f"E{r}"] = row.get("description"); ws[f"E{r}"].alignment = Alignment(wrap_text=True, vertical="top")
            ws[f"F{r}"] = row.get("source")
            ws[f"G{r}"] = conv(row.get("amount"), row.get("currency")); ws[f"G{r}"].number_format = money
            if row.get("notes"): ws[f"C{r}"] = row["notes"]
        elif kind == "landtotal":
            a, b = obj; frame(r, True)
            ws[f"E{r}"] = "Total cost"; ws[f"G{r}"] = f"=SUM(G{a}:G{b})"
            ws[f"H{r}"] = f"=G{r}*{round((land.get('markup') if land.get('markup') is not None else T.get('markup', 0.15))*100, 2):g}%"
            ws[f"J{r}"] = f"=G{r}+H{r}"; ws[f"J{r}"].font = F(b=True, color="FFFF0000")
            for col in "GHJ": ws[f"{col}{r}"].number_format = money

    # ── ROE block ──────────────────────────────────────────────────────────
    for c, rr in roe_row.items():
        e = ws[f"E{rr}"]; e.value = f"ROE - {c}"; e.font = F("Arial", True); e.border = ALLTHIN; e.alignment = Alignment(horizontal="left")
        v = ws[f"F{rr}"]; v.value = rates.get(c); v.number_format = ACC; v.border = ALLTHIN
        if rates.get(c) is None: v.fill = YELLOW
    r = roe_start + len(roe_list) + 4

    # ── Summary block ──────────────────────────────────────────────────────
    srow = r
    lab = ["Cost to US", "Markup ", "Cost to Client ", "GST", f"TCS {round(T.get('tcs', 0.02)*100, 2):g}%", "Payable", "Round Off"]
    for j, p in enumerate(pkgs):
        c1, c2 = CL(6 + 2*j), CL(7 + 2*j)
        for col, val in ((c1, p.get("name") or "Per couple"), (c2, "in INR")):
            cell = ws[f"{col}{srow}"]; cell.value = val; cell.font = F(b=True); cell.border = ALLTHIN; cell.alignment = Alignment(horizontal="center")
        rows_ = selected.get(p["key"], [])
        use_land = bool(land_total_row) and p["key"] in (land["pkg"] if "pkg" in land else [p["key"]])
        cost = "+".join([f"G{x}" for x in rows_] + ([f"G{land_total_row}"] if use_land else [])) or "0"
        mark = "+".join([f"H{x}" for x in rows_] + ([f"H{land_total_row}"] if use_land else [])) or "0"
        b = srow + 1
        fx = [f"={cost}", f"={mark}", f"=SUM({c1}{b}:{c1}{b+1})", f"={c1}{b+2}*{round(T.get('gst', 0.05)*100, 2):g}%",
              f"=({c1}{b+2}+{c1}{b+3})*{round(T.get('tcs', 0.02)*100, 2):g}%", f"={c1}{b+2}+{c1}{b+3}+{c1}{b+4}", None]
        for i, f_ in enumerate(fx):
            rr = b + i
            a = ws[f"{c1}{rr}"]; a.value = f_; a.font = F("Arial"); a.fill = BEIGE; a.border = ALLTHIN; a.number_format = money
            inr = ws[f"{c2}{rr}"]; inr.font = F("Arial"); inr.fill = BEIGE; inr.border = ALLTHIN; inr.number_format = FMT["INR"]
            if cc == "INR": inr.value = f"={c1}{rr}" if f_ else f"=ROUNDUP({c2}{rr-1},-2)"
            else: inr.value = f"={c1}{rr}*{ROE(cc)}" if f_ else f"=ROUNDUP({c2}{rr-1},-2)"
    for i, t in enumerate(lab):
        e = ws[f"E{srow+1+i}"]; e.value = t; e.font = F("Arial", True, "FF00B050" if t == "Round Off" else th(3)); e.fill = BEIGE; e.border = ALLTHIN
    r = srow + 1 + len(lab) + 3

    # ── Commission + Reporting blocks ──────────────────────────────────────
    def titled(r, title, heads, rows_):
        ws.merge_cells(f"E{r}:G{r}")
        t = ws[f"E{r}"]; t.value = title; t.font = F("Arial", True, NAVYTXT); t.fill = BLUE; t.alignment = Alignment(horizontal="center")
        for col in "EFG": ws[f"{col}{r}"].border = ALLTHIN
        for j, h in enumerate(heads):
            c = ws.cell(r + 1, 5 + j, h); c.font = F("Arial", True, th(3)); c.fill = BEIGE; c.border = ALLTHIN
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for i, vals in enumerate(rows_):
            for j in range(len(heads)):
                c = ws.cell(r + 2 + i, 5 + j, vals[j] if j < len(vals) else None); c.border = ALLTHIN
        return r + 2 + len(rows_)
    r = titled(r, "commission ", ["Hotel Name", "Check out date", "Amount"], [[]] * 6) + 2
    seen, rep = set(), []
    for kind, rr, obj in layout:
        if kind == "row" and obj.get("pkg") and obj.get("description") not in seen:
            seen.add(obj.get("description")); rep.append([obj.get("description"), obj.get("source")])
    titled(r, "Reporting", ["Hotel/DMC", "Source", ""], rep + [[]] * max(0, 8 - len(rep)))
    wb.calculation.fullCalcOnLoad = True
    wb.save(out)

if __name__ == "__main__":
    build(json.load(open(sys.argv[1], encoding="utf-8")), sys.argv[2]); print("saved", sys.argv[2])
